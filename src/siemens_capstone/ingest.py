"""Read workbook rows locally and expose aggregate quality evidence."""
import csv
from datetime import date, datetime, time, timedelta
from hashlib import sha256
from io import BytesIO, StringIO
import math
import re
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel, CALENDAR_WINDOWS_1900, CALENDAR_MAC_1904

SHEETS = {
    'HISTORICAL_PROD_DATA': ['salesorderOrdNr', 'productionorderSalesOrderLineItem', 'jobtraceObjid'],
    'PLANNED_PROD_SCHEDULE': ['Sales Order', 'Line Item', 'Basic Start Date'],
    'HISTORICAL_SM_REPLENISHMENTS': ['Transfer Order Number', 'Transfer order item', 'Material', 'Dest.Storage Bin', 'Dest.target quantity'],
    'BIN_STATUS_EXAMPLE': ['Material', 'Storage Bin', 'Available stock', 'Base Unit of Measure'],
}
KEYS = {
    'HISTORICAL_PROD_DATA': ['jobtraceObjid'],
    'PLANNED_PROD_SCHEDULE': ['Sales Order', 'Line Item'],
    'HISTORICAL_SM_REPLENISHMENTS': ['Transfer Order Number', 'Transfer order item'],
    'BIN_STATUS_EXAMPLE': ['Material', 'Storage Bin', 'Storage Unit', 'Special Stock Number'],
}


def excel_datetime(day, fraction=0, system='1900'):
    if system not in ('1900', '1904'):
        raise ValueError('Unknown workbook date system.')
    epoch = CALENDAR_WINDOWS_1900 if system == '1900' else CALENDAR_MAC_1904
    try:
        if isinstance(day, datetime):
            dt = day
        elif isinstance(day, date):
            dt = datetime.combine(day, time())
        elif isinstance(day, (int, float)) and math.isfinite(day):
            dt = from_excel(day, epoch)
            if isinstance(dt, time):
                dt = datetime.combine(epoch.date(), dt)
        else:
            dt = datetime.fromisoformat(str(day))
        if isinstance(fraction, time):
            seconds = fraction.hour * 3600 + fraction.minute * 60 + fraction.second
        elif isinstance(fraction, str) and ':' in fraction:
            t = time.fromisoformat(fraction); seconds = t.hour * 3600 + t.minute * 60 + t.second
        else:
            f = float(fraction or 0)
            if not math.isfinite(f) or not 0 <= f < 1:
                raise ValueError()
            seconds = f * 86400
        return (dt + timedelta(seconds=seconds)).isoformat()
    except (ValueError, TypeError, OverflowError):
        raise ValueError('Invalid Excel date or time.') from None


def identifier(value):
    s = str(value).strip() if value is not None else ''
    if re.fullmatch(r'\d+(\.0+)?', s):
        return str(int(s.split('.')[0]))
    return s


def serializable(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def inspect_workbook(content, filename):
    if not content or len(content) > 32 * 1024 * 1024:
        raise ValueError('Upload an XLSX workbook under 32 MB.')
    try:
        with ZipFile(BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 350 * 1024 * 1024:
                raise ValueError('Expanded workbook exceeds the local import limit.')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        # Parser errors must not leak local internals through the API.
        raise ValueError('The file is not a readable XLSX workbook or exceeds the import limit.') from exc
    findings, tables, summaries = [], {}, []
    system = '1904' if wb.epoch == CALENDAR_MAC_1904 else '1900'
    try:
        for name, required in SHEETS.items():
            if name not in wb.sheetnames:
                findings.append(dict(code='missing_sheet', severity='error', sheet=name, count=1, message='Required sheet is missing.'))
                tables[name] = []; continue
            ws = wb[name]
            if ws.max_row and ws.max_row > 300001:
                raise ValueError('A sheet exceeds the 300,000 row limit.')
            rows = ws.iter_rows(values_only=True)
            header = [str(v).strip() if v is not None else '' for v in next(rows, [])]
            if len([h for h in header if h]) != len(set(h for h in header if h)):
                raise ValueError(f'{name} has duplicate column headings.')
            missing = [k for k in required if k not in header]
            if missing:
                findings.append(dict(code='missing_column', severity='error', sheet=name, count=len(missing), message='Missing: '+', '.join(missing)))
            seen = set(); counters = dict(duplicate_key=0, missing_key=0, negative_quantity=0, zero_quantity=0, invalid_number=0, time_order=0, invalid_date=0, unsupported_unit=0)
            clean = []
            key_fields = [k for k in KEYS[name] if k in header]
            for index, raw in enumerate(rows):
                if index >= 300000:
                    raise ValueError('A sheet exceeds the 300,000 row limit.')
                if all(v is None for v in raw): continue
                row = {k: serializable(v) for k, v in zip(header, raw) if k}
                key = tuple(identifier(row.get(k)) for k in key_fields)
                essential = [k for k in key_fields if k not in ('Storage Unit', 'Special Stock Number')]
                if any(not identifier(row.get(k)) for k in essential): counters['missing_key'] += 1
                if key in seen: counters['duplicate_key'] += 1
                seen.add(key)
                for k in ('Available stock', 'Total Stock', 'Source target qty', 'Dest.target quantity'):
                    if k in row:
                        try:
                            n = float(row[k])
                            if not math.isfinite(n): raise ValueError()
                            if n < 0: counters['negative_quantity'] += 1
                            if n == 0 and k == 'Dest.target quantity': counters['zero_quantity'] += 1
                        except (TypeError, ValueError): counters['invalid_number'] += 1
                if name == 'BIN_STATUS_EXAMPLE' and row.get('Base Unit of Measure') not in ('PC', 'EA'):
                    counters['unsupported_unit'] += 1
                if name == 'HISTORICAL_SM_REPLENISHMENTS':
                    if row.get('Creation Date') is not None and row.get('Confirmation Date') is not None:
                        try:
                            start = excel_datetime(row['Creation Date'], row.get('Creation time', 0), system)
                            end = excel_datetime(row['Confirmation Date'], row.get('Confirmation time', 0), system)
                            row['_created_at_local'] = start; row['_confirmed_at_local'] = end
                            if datetime.fromisoformat(end) < datetime.fromisoformat(start): counters['time_order'] += 1
                        except ValueError: counters['invalid_date'] += 1
                if name == 'PLANNED_PROD_SCHEDULE':
                    try: row['_start_at_local'] = excel_datetime(row.get('Basic Start Date'), 0, system)
                    except ValueError: counters['invalid_date'] += 1
                clean.append(row)
            for code, count in counters.items():
                if count:
                    messages = dict(duplicate_key='Repeated candidate keys; confirm grain before joining.', missing_key='Rows lack a required join identifier.', negative_quantity='Negative quantities require movement or reversal semantics.', zero_quantity='Zero destination quantities need review.', invalid_number='Quantity is missing or not numeric.', time_order='Confirmation precedes creation.', invalid_date='Date or time could not be interpreted.', unsupported_unit='Unit requires an explicit conversion contract.')
                    findings.append(dict(code=code, severity='warning', sheet=name, count=count, message=messages[code]))
            tables[name] = clean
            summaries.append(dict(name=name, rows=len(clean), columns=len(header), fields=header, required_missing=missing,
                                  grain=' + '.join(key_fields)))
    finally:
        wb.close()
    joins = []
    prod = {(identifier(r.get('salesorderOrdNr')), identifier(r.get('productionorderSalesOrderLineItem'))) for r in tables.get('HISTORICAL_PROD_DATA', [])}
    planned = tables.get('PLANNED_PROD_SCHEDULE', [])
    matched = sum((identifier(r.get('Sales Order')), identifier(r.get('Line Item'))) in prod and bool(identifier(r.get('Sales Order'))) and bool(identifier(r.get('Line Item'))) for r in planned)
    joins.append(dict(name='Schedule → production', matched=matched, total=len(planned), description='Schedule rows with matching sales order + line item in production; event rows are reduced to distinct keys.'))
    bins = {(str(r.get('Material', '')).strip(), str(r.get('Storage Bin', '')).strip()) for r in tables.get('BIN_STATUS_EXAMPLE', [])}
    repl = tables.get('HISTORICAL_SM_REPLENISHMENTS', [])
    matched = sum((str(r.get('Material', '')).strip(), str(r.get('Dest.Storage Bin', '')).strip()) in bins and bool(r.get('Material')) and bool(r.get('Dest.Storage Bin')) for r in repl)
    joins.append(dict(name='Replenishments → bin snapshot', matched=matched, total=len(repl), description='Transfer rows with matching material + destination bin. These are possible links, not reconstructed consumption.'))
    blockers = ['No stock history yet: the bin snapshot is one undated extract, so depletion cannot be observed until dated snapshots are added.',
                'Warehouse lead time and the warning horizon are unconfirmed and need a business definition.',
                'The BOM is one example file; production-order and part coverage is partial and unvalidated.',
                'Min/Max tolerance gaps: positions without a tolerance row, zero minimums, and rows with min at or above max need business review.']
    return dict(filename=filename, checksum=sha256(content).hexdigest(), date_system=system, timezone='Unconfirmed',
                total_rows=sum(len(rows) for rows in tables.values()), sheets=summaries, findings=findings,
                joins=joins, blockers=blockers, forecast_ready=False, records=tables)


REFERENCE_KINDS = (  # (kind, required headers, retained columns), checked in this order
    ('bom', ('Item', 'Part', 'Qty'), ('Item', 'Part', 'Qty')),
    ('tolerances', ('Part', 'Location', 'Min', 'Max'), ('Part', 'Location', 'Type', 'Min', 'Max')),
    ('snapshot', ('Material', 'Storage Bin', 'Available stock'), ('Material', 'Storage Type', 'Storage Bin', 'Available stock', 'Base Unit of Measure', 'Storage Unit', 'Special Stock Number')),
)
REFERENCE_KEYS = dict(bom=None, tolerances=('Part', 'Location'), snapshot=('Material', 'Storage Type', 'Storage Bin', 'Storage Unit', 'Special Stock Number'))
REFERENCE_ESSENTIAL = dict(bom=('Item', 'Part'), tolerances=('Part', 'Location'), snapshot=('Material', 'Storage Bin'))
REFERENCE_NUMBERS = dict(bom=('Qty',), tolerances=('Min', 'Max'), snapshot=('Available stock',))


def is_full_workbook(content):
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
        try: return all(name in wb.sheetnames for name in SHEETS)
        finally: wb.close()
    except Exception:
        return False


def reference_tables(content, filename):
    if filename.lower().endswith('.csv'):
        try: text = content.decode('utf-8-sig')
        except UnicodeDecodeError: raise ValueError('The CSV file is not UTF-8 text.') from None
        yield 'csv', csv.reader(StringIO(text)); return
    try:
        with ZipFile(BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 350 * 1024 * 1024:
                raise ValueError('Expanded workbook exceeds the local import limit.')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except ValueError: raise
    except (BadZipFile, Exception) as exc:
        raise ValueError('The file is not a readable XLSX workbook or exceeds the import limit.') from exc
    try:
        for name in sorted(wb.sheetnames, key=lambda n: n != 'Full BOM'):
            if wb[name].max_row and wb[name].max_row > 300001: raise ValueError('A sheet exceeds the 300,000 row limit.')
            yield name, wb[name].iter_rows(values_only=True)
    finally:
        wb.close()


def inspect_reference(content, filename, as_of=None):
    """Classify a Min/Max tolerance, BOM, or dated stock-snapshot file and return aggregate evidence plus clean rows."""
    if not content or len(content) > 32 * 1024 * 1024:
        raise ValueError('Upload an XLSX or CSV file under 32 MB.')
    for sheet, rows in reference_tables(content, filename):
        rows = iter(rows)
        header = [str(v).strip() if v is not None else '' for v in next(rows, [])]
        kind = next((k for k, required, _ in REFERENCE_KINDS if all(h in header for h in required)), None)
        if kind is None: continue
        if len([h for h in header if h]) != len(set(h for h in header if h)):
            raise ValueError(f'{sheet} has duplicate column headings.')
        if kind == 'snapshot':
            try:
                if not isinstance(as_of, str) or len(as_of) != 10: raise ValueError()
                date.fromisoformat(as_of)
            except ValueError: raise ValueError('A stock snapshot needs a snapshot date (YYYY-MM-DD).') from None
        else: as_of = None
        keep = next(cols for k, _, cols in REFERENCE_KINDS if k == kind)
        key_fields = [k for k in (REFERENCE_KEYS[kind] or ()) if k in header]
        clean, seen, counters = [], set(), dict(duplicate_key=0, missing_key=0, invalid_number=0, negative_quantity=0)
        for index, raw in enumerate(rows):
            if index >= 300000: raise ValueError('A sheet exceeds the 300,000 row limit.')
            if all(v is None or v == '' for v in raw): continue
            row = {k: serializable(v) for k, v in zip(header, raw) if k in keep}
            key = tuple(identifier(row.get(k)) for k in key_fields)
            if any(not identifier(row.get(k)) for k in REFERENCE_ESSENTIAL[kind]): counters['missing_key'] += 1
            if key_fields:
                if key in seen: counters['duplicate_key'] += 1
                seen.add(key)
            for k in REFERENCE_NUMBERS[kind]:
                try:
                    n = float(row.get(k))
                    if not math.isfinite(n): raise ValueError()
                    if n < 0: counters['negative_quantity'] += 1
                except (TypeError, ValueError): counters['invalid_number'] += 1
            clean.append(row)
        messages = dict(duplicate_key='Repeated keys; the last row wins or quantities are added.', missing_key='Rows lack a required identifier.',
                        invalid_number='Quantity is missing or not numeric.', negative_quantity='Negative values need business review.')
        findings = [dict(code=c, severity='warning', sheet=sheet, count=n, message=messages[c]) for c, n in counters.items() if n]
        return dict(kind=kind, filename=filename, checksum=sha256(content).hexdigest(), rows=len(clean), findings=findings, as_of=as_of, records=clean)
    raise ValueError('Unrecognized reference file: expected Min/Max tolerances (Part, Location, Min, Max), a BOM (Item, Part, Qty), or a stock snapshot (Material, Storage Bin, Available stock).')
