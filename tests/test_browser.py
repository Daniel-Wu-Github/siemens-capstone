"""Real browser journeys against a temporary SQLite database."""
import json
from pathlib import Path

import pytest
playwright = pytest.importorskip("playwright")
from playwright.sync_api import sync_playwright, expect

from test_api import running_server
from test_ingest import workbook_bytes


@pytest.mark.browser
def test_dashboard_policy_delivery_import_evaluation_and_export(tmp_path):
    artifacts = Path('artifacts'); artifacts.mkdir(exist_ok=True)
    with running_server(tmp_path) as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1440, 'height': 1050}, device_scale_factor=1)
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        page.goto(base)
        expect(page.get_by_role('heading', name='Stay ahead of the next shortage.')).to_be_visible()
        expect(page.get_by_text('SYNTHETIC DEMO', exact=True)).to_be_visible()
        page.get_by_role('combobox', name='Chart time window').select_option('6')
        expect(page.get_by_text('Recent consumption and the next 6 hours', exact=True)).to_be_visible()
        page.get_by_role('combobox', name='Chart time window').select_option('12')
        page.screenshot(path=str(artifacts / 'overview.png'), full_page=True)
        page.get_by_role('button', name='Run forecast').click()
        expect(page.locator('#toast')).to_contain_text('Forecast and recommendations saved')
        page.get_by_role('button', name='Materials & alerts', exact=True).click()
        page.get_by_role('searchbox', name='Search materials').fill('DEM-101')
        expect(page.locator('#material-rows tr')).to_have_count(1)
        page.get_by_role('button', name='Inspect', exact=True).click()
        page.get_by_role('button', name='Edit policy', exact=True).click()
        page.get_by_label('Minimum stock (pieces)').fill('25')
        page.get_by_role('button', name='Save & recalculate').click()
        expect(page.get_by_role('dialog')).not_to_be_visible()
        page.reload()
        expect(page.get_by_role('heading', name='Stay ahead of the next shortage.')).to_be_visible()
        page.get_by_role('button', name='Edit policy', exact=True).click()
        expect(page.get_by_label('Minimum stock (pieces)')).to_have_value('25')
        page.get_by_role('button', name='Cancel', exact=True).click()

        # Dispatch does not change stock. A receipt changes it and survives reload.
        state = page.request.get(base + '/api/state?dataset=demo').json()
        alert = next(a for a in state['alerts'] if a['status'] == 'Open')
        old_stock = next(m['stock'] for m in state['materials'] if m['id'] == alert['material_id'])
        page.get_by_role('button', name='Materials & alerts', exact=True).click()
        page.locator(f'button[data-alert="{alert["id"]}"][data-next="Dispatched"]').click()
        expect(page.locator(f'button[data-deliver="{alert["id"]}"]')).to_be_visible()
        after_dispatch = page.request.get(base + '/api/state?dataset=demo').json()
        assert next(m['stock'] for m in after_dispatch['materials'] if m['id'] == alert['material_id']) == old_stock
        page.locator(f'button[data-deliver="{alert["id"]}"]').click()
        page.get_by_label('Quantity received (pieces)').fill('200')
        page.get_by_role('button', name='Confirm receipt', exact=True).click()
        expect(page.get_by_role('dialog')).not_to_be_visible()
        page.reload()
        after_delivery = page.request.get(base + '/api/state?dataset=demo').json()
        assert next(m['stock'] for m in after_delivery['materials'] if m['id'] == alert['material_id']) == old_stock + 200

        page.get_by_role('button', name='Model evaluation', exact=True).click()
        page.get_by_role('button', name='Run evaluation').click()
        expect(page.locator('#toast')).to_contain_text('Evaluation completed')
        expect(page.get_by_role('heading', name='Comparison report')).to_be_visible()
        page.screenshot(path=str(artifacts / 'evaluation.png'), full_page=True)
        with page.expect_download() as info:
            page.get_by_role('button', name='Export report', exact=False).click()
        downloaded = info.value
        report = json.loads(Path(downloaded.path()).read_text())
        assert report['evaluation']['sample_count'] == 60 and report['runs']

        page.get_by_role('button', name='Data workspace', exact=True).click()
        page.get_by_label('Import workbook').set_input_files({'name': 'synthetic-test.xlsx', 'mimeType': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'buffer': workbook_bytes()})
        expect(page.get_by_text('LOCAL WORKBOOK', exact=True)).to_be_visible()
        expect(page.get_by_text('negative quantity · 1 occurrences', exact=True)).to_be_visible()
        expect(page.get_by_text('Forecast readiness: definitions needed', exact=True)).to_be_visible()
        page.get_by_role('button', name='Model evaluation', exact=True).click()
        expect(page.get_by_role('button', name='Run evaluation')).to_be_disabled()
        page.get_by_role('combobox', name='Active dataset').select_option('demo')
        expect(page.get_by_role('heading', name='Comparison report')).to_be_visible()
        assert not errors
        browser.close()


@pytest.mark.browser
def test_mobile_navigation_and_scenario(tmp_path):
    with running_server(tmp_path) as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 390, 'height': 844}, is_mobile=True, has_touch=True)
        page.goto(base)
        expect(page.get_by_role('heading', name='Stay ahead of the next shortage.')).to_be_visible()
        page.locator('#demand').focus()
        for _ in range(10): page.locator('#demand').press('ArrowRight')
        page.get_by_role('button', name='Apply scenario & save').click()
        expect(page.locator('#toast')).to_contain_text('Production scenario saved')
        page.reload()
        expect(page.locator('#demand')).to_have_value('1.5')
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        page.screenshot(path='artifacts/mobile.png', full_page=True)
        browser.close()


@pytest.mark.browser
def test_workbook_insights_scenario_saved_and_source_preserved(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(workbook_bytes()))
    ws = wb['HISTORICAL_SM_REPLENISHMENTS']; ws['E2'] = 60; ws['H2'] = ws['F2'].value
    buf = BytesIO(); wb.save(buf)
    with running_server(tmp_path) as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1440, 'height': 1000})
        result = page.request.post(base + '/api/import?filename=synthetic-insights.xlsx',
                                   data=buf.getvalue(), headers={'Content-Type': 'application/octet-stream'})
        ds = result.json()['id']
        page.goto(base + '/?dataset=' + ds + '&view=insights')
        expect(page.get_by_role('heading', name='Your data, put to work.')).to_be_visible()
        expect(page.get_by_role('heading', name='Stock snapshot review')).to_be_visible()
        page.get_by_label('Assumed usage multiplier').fill('2')
        page.get_by_role('button', name='Calculate & save scenario').click()
        expect(page.locator('#workbook-scenario-result')).to_contain_text('2h')
        page.reload()
        expect(page.locator('#workbook-scenario-result')).to_contain_text('2h')
        page.get_by_label('Search workbook inventory').fill('no-such-part')
        expect(page.locator('#workbook-inventory')).to_contain_text('No positions match')
        page.get_by_label('Search workbook inventory').fill('DEMO-X')
        expect(page.locator('#workbook-inventory tr')).to_have_count(1)
        saved = page.request.get(base + '/api/insights?dataset=' + ds).json()
        assert saved['inventory'][0]['stock'] == 10
        assert saved['scenarios'][0]['kind'] == 'assumption_based'
        assert page.request.get(base + '/api/state?dataset=' + ds).json()['alerts'] == []
        browser.close()


@pytest.mark.browser
def test_scenario_response_after_dataset_switch_preserves_ui_context(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(workbook_bytes()))
    ws = wb['HISTORICAL_SM_REPLENISHMENTS']; ws['E2'] = 60; ws['H2'] = ws['F2'].value
    buf = BytesIO(); wb.save(buf)
    with running_server(tmp_path) as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        result = page.request.post(base + '/api/import?filename=synthetic-race.xlsx',
                                   data=buf.getvalue(), headers={'Content-Type': 'application/octet-stream'})
        ds = result.json()['id']
        page.goto(base + '/?dataset=' + ds + '&view=insights')
        expect(page.get_by_role('heading', name='Stock snapshot review')).to_be_visible()
        def switch_dataset(route):
            page.get_by_role('combobox', name='Active dataset').select_option('demo')
            expect(page.get_by_text('Select a manufacturing workbook', exact=True)).to_be_visible()
            route.continue_()
        page.route('**/api/workbook-scenario', switch_dataset)
        with page.expect_response('**/api/workbook-scenario'):
            page.get_by_role('button', name='Calculate & save scenario').click()
        expect(page.get_by_role('combobox', name='Active dataset')).to_have_value('demo')
        expect(page.get_by_text('Select a manufacturing workbook', exact=True)).to_be_visible()
        expect(page.locator('#toast')).to_contain_text('scenario saved')
        saved = page.request.get(base + '/api/insights?dataset=' + ds).json()
        assert len(saved['scenarios']) == 1
        assert page.locator('#workbook-scenario-result').count() == 0
        assert not errors
        browser.close()


@pytest.mark.browser
def test_reference_tolerances_upload_and_insights_status(tmp_path):
    from io import BytesIO
    from openpyxl import Workbook, load_workbook
    wb = load_workbook(BytesIO(workbook_bytes()))
    ws = wb['HISTORICAL_SM_REPLENISHMENTS']; ws['E2'] = 60; ws['H2'] = ws['F2'].value
    buf = BytesIO(); wb.save(buf)
    ref = Workbook(); sheet = ref.active; sheet.title = 'Sheet3'
    for row in [['Part', 'Description', 'Type', 'Location', 'Min', 'Max'], ['DEMO-X', 'Synthetic part', '401', 'A-1', 20, 100]]:
        sheet.append(row)
    ref_buf = BytesIO(); ref.save(ref_buf)
    with running_server(tmp_path) as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1440, 'height': 1000})
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        ds = page.request.post(base + '/api/import?filename=synthetic-ref.xlsx', data=buf.getvalue(),
                               headers={'Content-Type': 'application/octet-stream'}).json()['id']
        # Without reference files the insights view degrades to hints, not errors.
        page.goto(base + '/?dataset=' + ds + '&view=insights')
        expect(page.get_by_role('heading', name='Stock snapshot review')).to_be_visible()
        expect(page.get_by_text('Import the Min/Max file in Data workspace').first).to_be_visible()
        page.get_by_role('button', name='Data workspace', exact=True).click()
        expect(page.get_by_role('heading', name='Reference files')).to_be_visible()
        page.get_by_label('Import reference file').set_input_files({'name': 'tolerances.xlsx', 'mimeType': 'application/octet-stream', 'buffer': ref_buf.getvalue()})
        expect(page.locator('#toast')).to_contain_text('Reference file imported')
        expect(page.get_by_text('Min/Max tolerances · Imported')).to_be_visible()
        page.get_by_role('button', name='Workbook insights', exact=True).click()
        expect(page.locator('#workbook-inventory')).to_contain_text('Below min')
        expect(page.locator('#tile-below')).to_have_text('1')
        page.get_by_label('Warning horizon (hours)').fill('24')
        page.get_by_label('Workbook inventory filter').select_option('below')
        expect(page.locator('#workbook-inventory tr')).to_have_count(1)
        expect(page.get_by_label('Assumed minimum stock (from Min/Max when imported)')).to_have_value('20')
        expect(page.get_by_label('Warning horizon (h)')).to_be_visible()
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        assert not errors
        browser.close()


@pytest.mark.browser
def test_stock_alerts_page_edit_sends_email_and_reset_restores(tmp_path, monkeypatch):
    from test_alerts import BIN, TOL
    from test_ingest import table_bytes
    for k in ('SMTP_HOST', 'TEAMS_WEBHOOK_URL', 'ALERT_WEBHOOK_URL', 'ALERT_EMAIL_TO'): monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv('ALERT_OUTBOX', str(tmp_path / 'outbox'))
    wb = table_bytes({'HISTORICAL_PROD_DATA': [['salesorderOrdNr', 'productionorderSalesOrderLineItem', 'jobtraceObjid']],
                      'PLANNED_PROD_SCHEDULE': [['Sales Order', 'Line Item', 'Basic Start Date']],
                      'HISTORICAL_SM_REPLENISHMENTS': [['Transfer Order Number', 'Transfer order item', 'Material', 'Dest.Storage Bin', 'Dest.target quantity']],
                      'BIN_STATUS_EXAMPLE': BIN})
    with running_server(tmp_path) as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1440, 'height': 1000})
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        ds = page.request.post(base + '/api/import?filename=s.xlsx', data=wb, headers={'Content-Type': 'application/octet-stream'}).json()['id']
        page.request.post(base + f'/api/import?filename=t.xlsx&dataset={ds}', data=table_bytes({'Sheet3': TOL}), headers={'Content-Type': 'application/octet-stream'})
        page.goto(base + '/?dataset=' + ds + '&view=alerts')
        expect(page.get_by_role('heading', name='Parts, minimums, and available stock')).to_be_visible()
        expect(page.locator('#stock-alert-count')).to_have_text('3')
        box = page.get_by_label('Available stock for PART-A A1')
        box.fill('2'); box.press('Enter')
        expect(page.locator('#toast')).to_contain_text('email saved')
        expect(page.locator('#stock-alert-count')).to_have_text('4')
        expect(page.locator('#mail-preview')).to_contain_text('PART-A')
        assert len(list((tmp_path / 'outbox').glob('*.eml'))) == 2  # baseline announcement + this edit
        page.get_by_role('button', name='Reset simulated stock').click()
        expect(page.locator('#stock-alert-count')).to_have_text('4')  # first click only arms the button
        page.get_by_role('button', name='Click again to clear all edits').click()
        expect(page.locator('#stock-alert-count')).to_have_text('3')
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
        assert not errors
        browser.close()


def test_stock_alerts_refresh_after_reference_upload_and_digest_needs_confirmation(tmp_path, monkeypatch):
    from test_alerts import BIN, TOL
    from test_ingest import table_bytes
    for k in ('SMTP_HOST', 'TEAMS_WEBHOOK_URL', 'ALERT_WEBHOOK_URL', 'ALERT_EMAIL_TO'): monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv('ALERT_OUTBOX', str(tmp_path / 'outbox'))
    wb = table_bytes({'HISTORICAL_PROD_DATA': [['salesorderOrdNr', 'productionorderSalesOrderLineItem', 'jobtraceObjid']],
                      'PLANNED_PROD_SCHEDULE': [['Sales Order', 'Line Item', 'Basic Start Date']],
                      'HISTORICAL_SM_REPLENISHMENTS': [['Transfer Order Number', 'Transfer order item', 'Material', 'Dest.Storage Bin', 'Dest.target quantity']],
                      'BIN_STATUS_EXAMPLE': BIN})
    looser = tmp_path / 'looser.xlsx'
    looser.write_bytes(table_bytes({'Sheet3': [TOL[0]] + [[r[0], r[1], r[2], r[3], 0, 50] for r in TOL[1:]]}))  # every min 0 -> only empty PART-B alerts
    with running_server(tmp_path) as base, sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={'width': 1440, 'height': 1000})
        errors = []
        page.on('pageerror', lambda error: errors.append(str(error)))
        ds = page.request.post(base + '/api/import?filename=s.xlsx', data=wb, headers={'Content-Type': 'application/octet-stream'}).json()['id']
        page.request.post(base + f'/api/import?filename=t.xlsx&dataset={ds}', data=table_bytes({'Sheet3': TOL}), headers={'Content-Type': 'application/octet-stream'})
        page.goto(base + '/?dataset=' + ds + '&view=alerts')
        expect(page.locator('#stock-alert-count')).to_have_text('3')
        mails = lambda: len(list((tmp_path / 'outbox').glob('*.eml')))
        before = mails()
        page.get_by_role('button', name='✉ Send digest').click()
        page.wait_for_timeout(300)
        assert mails() == before  # first click only arms the button
        page.get_by_role('button', name='Click again to send').click()
        expect(page.locator('#toast')).to_contain_text('Digest')
        assert mails() == before + 1
        page.get_by_role('button', name='Data workspace', exact=True).click()
        page.locator('#ref-input').set_input_files(str(looser))
        expect(page.locator('#toast')).to_contain_text('Reference file imported')
        page.get_by_role('button', name='Stock alerts', exact=True).click()
        expect(page.locator('#stock-alert-count')).to_have_text('1')
        assert not errors
        browser.close()
