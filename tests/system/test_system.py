"""Table A7.3: system tests ST-01 to ST-07.

These drive the whole prototype: the store, the loopback server, and the
dashboard in a real browser.
"""
from http.server import ThreadingHTTPServer
import re
import subprocess
import threading
import time

import pytest

from siemens_capstone import alerts, pipeline
from siemens_capstone.api import handler_for
from siemens_capstone.store import Store

import a11y
from conftest import TOLERANCES, WORKBOOK, requires_real_data

pytestmark = pytest.mark.system

PIPELINE_BUDGET_SECONDS = 300    # ST-02: under 5 minutes
DASHBOARD_BUDGET_SECONDS = 3     # ST-03: under 3 seconds
ACCESSIBILITY_TARGET = 90        # ST-07


@pytest.fixture(scope='module')
def live_store(tmp_path_factory):
    """One store holding the real workbook and tolerances, shared by ST-01 and ST-03."""
    from conftest import MISSING, NEEDS_DATA
    if MISSING:
        pytest.skip(NEEDS_DATA)
    work = tmp_path_factory.mktemp('system')
    config = alerts.config({'ALERT_OUTBOX': str(work / 'outbox'),
                            'ALERT_DASHBOARD_URL': 'http://127.0.0.1:8765/#stock'})
    store = Store(work / 'capstone.sqlite3', alert_config=config)
    dataset = store.import_workbook(WORKBOOK.read_bytes(), WORKBOOK.name)
    store.import_reference(dataset['id'], TOLERANCES.read_bytes(), TOLERANCES.name)
    return store, dataset['id'], work


@pytest.fixture(scope='module')
def live_server(live_store):
    """The real loopback server, on a port the operating system picks."""
    store, dataset_id, _ = live_store
    server = ThreadingHTTPServer(('127.0.0.1', 0), handler_for(store))
    server.daemon_threads = True
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f'http://127.0.0.1:{server.server_port}', dataset_id
    finally:
        server.shutdown()
        server.server_close()


@pytest.fixture(scope='module')
def browser_page(live_server):
    page_module = pytest.importorskip('playwright.sync_api',
                                      reason='needs playwright: pip install -r requirements-dev.txt')
    try:
        runner = page_module.sync_playwright().start()
    except Exception as exc:                                    # pragma: no cover
        pytest.skip(f'playwright could not start: {exc}')
    try:
        browser = runner.chromium.launch()
    except Exception as exc:
        runner.stop()
        pytest.skip(f'needs chromium: .venv/bin/python -m playwright install chromium ({exc})')
    page = browser.new_page(viewport={'width': 1440, 'height': 900})
    try:
        yield page
    finally:
        page.close()
        browser.close()
        runner.stop()


def open_stock_view(page, base_url, dataset_id):
    """Open the dashboard on the imported workbook's Stock alerts page."""
    page.goto(base_url, wait_until='domcontentloaded')
    # Options are never "visible" to Playwright, so wait on the list being populated.
    page.wait_for_function(
        "id => [...document.querySelectorAll('#dataset option')].some(o => o.value === id)",
        arg=dataset_id, timeout=30000)
    page.select_option('#dataset', dataset_id)
    page.click('button[data-view="alerts"]')
    page.wait_for_selector('#stock-rows tr', timeout=30000)


@requires_real_data()
def test_ST01_full_daily_run(live_store):
    """Input: all data plus one simulated daily snapshot.
    Expected: dashboard updates and alert is sent with no errors."""
    store, dataset_id, _ = live_store

    before = store.stock_view(dataset_id)
    assert before['summary']['positions'] > 0
    assert before['tolerances_loaded'] is True

    # One simulated daily snapshot reading: drive a healthy bin below its minimum.
    healthy = next(p for p in before['positions']
                   if p['status'] == 'OK' and p['min'] is not None and p['min'] > 0)
    result = store.set_available(dataset_id, healthy['key'], 0)

    assert result['new'] == 1, 'the new reading should open exactly one alert'
    notification = result['notification']
    assert notification is not None, 'an alert must be sent for the new breach'
    assert [r for r in notification['results'] if r['status'] in ('sent', 'saved')], \
        f'no channel accepted the alert: {notification["results"]}'

    after = store.stock_view(dataset_id)
    assert after['summary']['alerts'] == before['summary']['alerts'] + 1
    assert any(a['position'] == healthy['key'] for a in after['open_alerts'])

    store.set_available(dataset_id, healthy['key'], None)   # leave the dataset as found


@requires_real_data()
def test_ST02_pipeline_speed(real_model, real_cleaned, parse_timings, config):
    """Input: full May to September data. Expected: finishes in under 5 minutes."""
    started = time.perf_counter()
    rows = pipeline.alert_rows(real_model, config['lead_hours'], config['buffer_hours'])
    view = pipeline.dashboard_view(real_model, config['lead_hours'], config['buffer_hours'])
    parse_timings['alert_and_dashboard'] = time.perf_counter() - started

    total = sum(parse_timings.values())
    stages = ', '.join(f'{k} {v:.1f}s' for k, v in sorted(parse_timings.items()))
    assert total < PIPELINE_BUDGET_SECONDS, \
        f'pipeline took {total:.1f}s, over the {PIPELINE_BUDGET_SECONDS}s budget ({stages})'

    # The run must actually have processed the full extract.
    assert real_cleaned['rows_in'] > 100_000, 'the full extract was not processed'
    assert view['summary']['positions'] > 0 and isinstance(rows, list)
    print(f'\nST-02 pipeline {total:.1f}s over {real_cleaned["rows_in"]:,} rows ({stages})')


@requires_real_data()
@pytest.mark.browser
def test_ST03_dashboard_speed(browser_page, live_server):
    """Input: open the dashboard and filter by supermarket.
    Expected: loads in under 3 seconds."""
    base_url, dataset_id = live_server

    started = time.perf_counter()
    open_stock_view(browser_page, base_url, dataset_id)
    load_seconds = time.perf_counter() - started

    first = browser_page.text_content('#stock-count') or ''

    started = time.perf_counter()
    browser_page.fill('#stock-search', 'Yellow')
    browser_page.wait_for_function(
        "prev => { const el = document.querySelector('#stock-count'); return el && el.textContent !== prev; }",
        arg=first, timeout=10000)
    filter_seconds = time.perf_counter() - started

    rows = browser_page.locator('#stock-rows tr').count()
    print(f'\nST-03 load {load_seconds:.2f}s, supermarket filter {filter_seconds:.2f}s, {rows} rows shown')

    assert load_seconds < DASHBOARD_BUDGET_SECONDS, f'dashboard took {load_seconds:.2f}s to load'
    assert filter_seconds < DASHBOARD_BUDGET_SECONDS, f'filter took {filter_seconds:.2f}s'
    assert rows > 0, 'filtering by supermarket returned nothing'


DATA_SUFFIXES = ('.xlsx', '.xls', '.csv', '.tsv', '.parquet', '.feather',
                 '.db', '.sqlite3', '.pkl', '.pdf')


def test_ST04_privacy_check(repo_root):
    """Input: all output files and the GitHub repo.
    Expected: no customer names or addresses and no Siemens files."""
    tracked = subprocess.run(['git', 'ls-files'], cwd=repo_root, capture_output=True,
                             text=True, check=True).stdout.split()

    committed_data = [f for f in tracked if f.lower().endswith(DATA_SUFFIXES)]
    assert committed_data == [], f'Siemens-shaped data files are committed: {committed_data}'

    # The folders holding real data and generated output must stay ignored.
    for folder in ('data/', 'outputs/', 'artifacts/'):
        ignored = subprocess.run(['git', 'check-ignore', '-q', folder], cwd=repo_root)
        assert ignored.returncode == 0, f'{folder} is not git-ignored'

    if not WORKBOOK.exists():
        pytest.skip('needs data/bin_status.xlsx to check for leaked identifiers')

    # Real identifiers taken from the supplied workbook must appear in no tracked file.
    from openpyxl import load_workbook
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        rows = book['HISTORICAL_PROD_DATA'].iter_rows(values_only=True)
        header = list(next(rows))
        columns = [header.index(c) for c in ('salesorderOrdNr', 'productionorderOrdNr') if c in header]
        secrets = set()
        for index, row in enumerate(rows):
            if index >= 400:
                break
            for column in columns:
                value = str(row[column]).strip() if row[column] is not None else ''
                if len(value) >= 6 and re.fullmatch(r'[\w.-]+', value):
                    secrets.add(value.removesuffix('.0'))
    finally:
        book.close()
    assert secrets, 'no identifiers were sampled, so the check would be vacuous'

    leaks = []
    for name in tracked:
        path = repo_root / name
        if not path.is_file() or path.stat().st_size > 4_000_000:
            continue
        try:
            body = path.read_text(errors='ignore')
        except OSError:
            continue
        for secret in secrets:
            if re.search(rf'\b{re.escape(secret)}\b', body):
                leaks.append(f'{name} contains {secret}')
    assert leaks == [], f'customer identifiers are committed: {leaks[:10]}'


def test_ST05_missing_file(tmp_path):
    """Input: run without the production schedule file.
    Expected: stops with a clear error message naming the file."""
    records = {'HISTORICAL_PROD_DATA': [{'jobtraceObjid': 'J1'}],
               'HISTORICAL_SM_REPLENISHMENTS': [{'Material': 'A'}],
               'BIN_STATUS_EXAMPLE': [{'Material': 'A', 'Storage Bin': 'B'}]}

    with pytest.raises(ValueError) as caught:
        pipeline.require_sources(records, filename='bin_status.xlsx')

    message = str(caught.value)
    assert 'PLANNED_PROD_SCHEDULE' in message, 'the error must name the missing source'
    assert 'production schedule' in message
    assert 'bin_status.xlsx' in message, 'the error must name the file it looked in'

    # A complete set runs without complaint.
    records['PLANNED_PROD_SCHEDULE'] = [{'Sales Order': '1', 'Line Item': '10'}]
    assert pipeline.require_sources(records, 'bin_status.xlsx') is True


def test_ST06_change_lead_time_in_config(tmp_path):
    """Input: change lead time from 3 to 4 hours in the config file.
    Expected: warning list updates without changing code."""
    # One bin that reaches its minimum in 3.5 hours: inside a 4h lead, outside a 3h lead.
    part = dict(part='PART-EDGE', bin='Y-99', supermarket='Yellow', available=55, min=20,
                max=100, below_minimum=False, hours_to_minimum=3.5, suggested_quantity=45)
    model = dict(parts=[part], modelled=1, not_enough_history=0, window_days=30)

    def warnings_for(lead_hours):
        path = tmp_path / 'minalert.local.toml'
        path.write_text(f'lead_hours = {lead_hours}\nbuffer_hours = 0\nwarn_days = 7\n')
        settings = pipeline.load_config(path)
        assert settings['lead_hours'] == lead_hours
        return pipeline.alert_rows(model, settings['lead_hours'], settings['buffer_hours'])

    assert warnings_for(3) == [], 'at a 3 hour lead time the bin is not yet a warning'

    raised = warnings_for(4)
    assert len(raised) == 1, 'at a 4 hour lead time the same bin must warn'
    assert raised[0]['part'] == 'PART-EDGE'

    # The config file is the only thing that changed.
    assert pipeline.DEFAULT_CONFIG['lead_hours'] == 3.0


@requires_real_data()
@pytest.mark.browser
def test_ST07_accessibility(browser_page, live_server):
    """Input: run a Lighthouse audit on the dashboard.
    Expected: score of 90 or higher."""
    base_url, dataset_id = live_server
    open_stock_view(browser_page, base_url, dataset_id)

    if a11y.lighthouse_available():
        result = a11y.run_lighthouse(base_url)
    else:
        result = a11y.run_local_audit(browser_page)

    failures = a11y.failed_audits(result)
    detail = '; '.join(f"{a['id']}: {a.get('failures') or a.get('title')}" for a in failures)
    print(f"\nST-07 accessibility {result['score']}/100 via {result['tool']}"
          + (f' — failing: {detail}' if detail else ' — all checks passed'))

    assert result['score'] >= ACCESSIBILITY_TARGET, \
        f"accessibility scored {result['score']}, below {ACCESSIBILITY_TARGET}. Failing: {detail}"
