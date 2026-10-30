from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import Workbook

from siemens_capstone.ingest import inspect_workbook, excel_datetime
from siemens_capstone.store import Store


def workbook_bytes():
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in {
        'HISTORICAL_PROD_DATA': [['salesorderOrdNr', 'productionorderSalesOrderLineItem', 'jobtraceObjid'], ['SO1', '0010', 'J1']],
        'PLANNED_PROD_SCHEDULE': [['Sales Order', 'Line Item', 'Basic Start Date'], ['SO1', '0010', datetime(2026, 9, 29)], ['SO2', '20', datetime(2026, 10, 1)]],
        'HISTORICAL_SM_REPLENISHMENTS': [['Transfer Order Number', 'Transfer order item', 'Material', 'Dest.Storage Bin', 'Dest.target quantity', 'Creation Date', 'Creation time', 'Confirmation Date', 'Confirmation time'], ['T1', '1', 'DEMO-X', 'A-1', -2, datetime(2026, 9, 29), 0.5, datetime(2026, 9, 28), 0.5]],
        'BIN_STATUS_EXAMPLE': [['Material', 'Storage Bin', 'Available stock', 'Base Unit of Measure'], ['DEMO-X', 'A-1', 10, 'PC'], ['DEMO-X', 'A-1', 10, 'PC']],
    }.items():
        ws = wb.create_sheet(name)
        for row in rows: ws.append(row)
    out = BytesIO(); wb.save(out); return out.getvalue()


def test_excel_date_systems_and_fractional_time():
    assert excel_datetime(1, 0.5, '1900') == '1900-01-01T12:00:00'
    assert excel_datetime(1, 0, '1904') == '1904-01-02T00:00:00'
    with pytest.raises(ValueError): excel_datetime('invalid', 0, '1900')


def test_inspection_reports_duplicates_bad_quantity_and_join_gaps():
    result = inspect_workbook(workbook_bytes(), 'synthetic.xlsx')
    assert result['total_rows'] == 6
    assert result['forecast_ready'] is False
    assert any(x['code'] == 'duplicate_key' for x in result['findings'])
    assert any(x['code'] == 'negative_quantity' for x in result['findings'])
    assert any(x['code'] == 'time_order' for x in result['findings'])
    join = next(j for j in result['joins'] if j['name'] == 'Schedule → production')
    assert join['matched'] == 1 and join['total'] == 2


def test_import_deduplicated_by_checksum_persists_rows(tmp_path):
    store = Store(tmp_path / 'test.sqlite3')
    first = store.import_workbook(workbook_bytes(), 'synthetic.xlsx')
    second = store.import_workbook(workbook_bytes(), 'synthetic.xlsx')
    assert first['id'] == second['id']
    sources = Store(tmp_path / 'test.sqlite3').sources(first['id'])
    assert sources['dataset']['kind'] == 'workbook'
    assert sources['quality']['total_rows'] == 6
    assert store.record_count(first['id']) == 6


def test_missing_sheets_and_invalid_file_do_not_make_a_dataset(tmp_path):
    store = Store(tmp_path / 'test.sqlite3')
    before = len(store.datasets())
    with pytest.raises(ValueError): store.import_workbook(b'not an excel file', 'bad.xlsx')
    assert len(store.datasets()) == before


def table_bytes(sheets):
    wb = Workbook(); wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows: ws.append(row)
    out = BytesIO(); wb.save(out); return out.getvalue()


TOLERANCES = [['Part', 'Description', 'Type', 'Location', 'Min', 'Max'], ['PART-A', 'Synthetic', '401', 'A1', 5, 20], ['PART-B', 'Synthetic', '401', 'B1', 0, 0]]
BOM = [['Item', 'Description', 'Part', 'Qty', 'Mat Type', 'SubAssembly'], [10, 'Synthetic', 'PART-A', 2, 'X', ''], [10, 'Synthetic', 'US2:PART-C', 1, 'X', '']]
SNAPSHOT = [['Material', 'Storage Bin', 'Available stock'], ['PART-A', 'A1', 8]]


def test_inspect_reference_detects_kind_by_headers():
    from siemens_capstone.ingest import inspect_reference
    tol = inspect_reference(table_bytes({'Sheet3': TOLERANCES}), 't.xlsx')
    assert tol['kind'] == 'tolerances' and tol['rows'] == 2 and tol['as_of'] is None and tol['checksum']
    bom = inspect_reference(table_bytes({'Routings': [['Item', 'Routing'], [10, 'R']], 'Full BOM': BOM}), 'b.xlsx')
    assert bom['kind'] == 'bom' and bom['rows'] == 2
    snap = inspect_reference(table_bytes({'S': SNAPSHOT}), 's.xlsx', '2026-10-01')
    assert snap['kind'] == 'snapshot' and snap['as_of'] == '2026-10-01'
    csv = inspect_reference(b'Material,Storage Bin,Available stock\nPART-A,A1,4\n', 's.csv', '2026-10-02')
    assert csv['kind'] == 'snapshot' and csv['rows'] == 1


def test_inspect_reference_rejects_unknown_undated_and_oversized():
    from siemens_capstone.ingest import inspect_reference
    with pytest.raises(ValueError, match='Unrecognized reference file'): inspect_reference(table_bytes({'S': [['A', 'B'], [1, 2]]}), 'x.xlsx')
    with pytest.raises(ValueError, match='snapshot date'): inspect_reference(table_bytes({'S': SNAPSHOT}), 's.xlsx')
    with pytest.raises(ValueError, match='snapshot date'): inspect_reference(table_bytes({'S': SNAPSHOT}), 's.xlsx', '10/01/2026')
    with pytest.raises(ValueError): inspect_reference(b'', 'x.xlsx')
    with pytest.raises(ValueError): inspect_reference(b'not a workbook', 'x.xlsx')


def test_inspect_reference_reports_duplicates_and_bad_numbers():
    from siemens_capstone.ingest import inspect_reference
    rows = TOLERANCES + [['PART-A', 'Dup', '401', 'A1', 'x', 9]]
    codes = {f['code'] for f in inspect_reference(table_bytes({'Sheet3': rows}), 't.xlsx')['findings']}
    assert {'duplicate_key', 'invalid_number'} <= codes


def test_blockers_describe_current_gaps():
    blockers = ' '.join(inspect_workbook(workbook_bytes(), 's.xlsx')['blockers'])
    assert 'stock history' in blockers and 'lead time' in blockers and 'BOM' in blockers and 'tolerance' in blockers.lower()
