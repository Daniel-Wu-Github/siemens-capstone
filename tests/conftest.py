"""Shared fixtures. Real Siemens workbooks stay in data/ and are never committed."""
from datetime import date
from pathlib import Path
import sys
import time

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
sys.path.insert(0, str(Path(__file__).parent))

DATA = ROOT / 'data'
WORKBOOK = DATA / 'bin_status.xlsx'          # four source sheets
TOLERANCES = DATA / 'min_max.xlsx'           # Min/Max tolerance list
SPONSOR = DATA / 'supermarket_summary.xlsx'  # table the sponsor sent on 2026-10-01

REAL_FILES = {'workbook': WORKBOOK, 'tolerances': TOLERANCES, 'sponsor': SPONSOR}
MISSING = [name for name, path in REAL_FILES.items() if not path.exists()]
NEEDS_DATA = f"needs the Siemens files in data/ (missing: {', '.join(MISSING) or 'none'})"

# The supplied extract covers May to September 2026.
TRAIN_WINDOW = (date(2026, 5, 1), date(2026, 7, 31))
TEST_WINDOW = (date(2026, 8, 1), date(2026, 9, 30))


def requires_real_data():
    return pytest.mark.skipif(bool(MISSING), reason=NEEDS_DATA)


@pytest.fixture(scope='session')
def repo_root():
    return ROOT


@pytest.fixture(scope='session')
def parse_timings():
    """Elapsed seconds for each stage, filled in by the fixtures below (used by ST-02)."""
    return {}


@pytest.fixture(scope='session')
def real_report(parse_timings):
    """Parse the full workbook once; every real-data test reuses this."""
    if MISSING:
        pytest.skip(NEEDS_DATA)
    from siemens_capstone.ingest import inspect_workbook
    started = time.perf_counter()
    report = inspect_workbook(WORKBOOK.read_bytes(), WORKBOOK.name)
    parse_timings['parse_workbook'] = time.perf_counter() - started
    return report


@pytest.fixture(scope='session')
def real_records(real_report):
    return real_report['records']


@pytest.fixture(scope='session')
def real_tolerances(parse_timings):
    if MISSING:
        pytest.skip(NEEDS_DATA)
    from siemens_capstone.ingest import inspect_reference
    started = time.perf_counter()
    report = inspect_reference(TOLERANCES.read_bytes(), TOLERANCES.name)
    parse_timings['parse_tolerances'] = time.perf_counter() - started
    assert report['kind'] == 'tolerances'
    return report['records']


@pytest.fixture(scope='session')
def sponsor_summary():
    """The sponsor's supermarket table, read straight from the supplied file."""
    if MISSING:
        pytest.skip(NEEDS_DATA)
    from openpyxl import load_workbook
    book = load_workbook(SPONSOR, read_only=True, data_only=True)
    try:
        sheet = book[book.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = [str(v).strip() if v is not None else '' for v in next(rows)]
        return [dict(zip(header, row)) for row in rows if any(v is not None for v in row)]
    finally:
        book.close()


@pytest.fixture(scope='session')
def real_events(real_records):
    from siemens_capstone import pipeline
    return pipeline.replenishment_events(real_records)


@pytest.fixture(scope='session')
def real_cleaned(real_records, real_tolerances, parse_timings):
    from siemens_capstone import pipeline
    started = time.perf_counter()
    cleaned = pipeline.clean_all(real_records, real_tolerances)
    parse_timings['clean_all'] = time.perf_counter() - started
    return cleaned


@pytest.fixture(scope='session')
def real_model(real_cleaned, real_tolerances, parse_timings):
    from siemens_capstone import pipeline
    started = time.perf_counter()
    model = pipeline.build_model(real_cleaned, real_tolerances)
    parse_timings['build_model'] = time.perf_counter() - started
    return model


@pytest.fixture(scope='session')
def config():
    """Operational settings, from the local TOML file when it exists."""
    from siemens_capstone import pipeline
    local = ROOT / 'minalert.local.toml'
    return pipeline.load_config(local if local.exists() else None)
