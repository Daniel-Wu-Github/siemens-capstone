from copy import deepcopy

import pytest

from siemens_capstone.analytics import analyze, scenario
from siemens_capstone.store import Store


def records():
    first = {'Transfer Order Number': 'T1', 'Transfer order item': '1', 'Material': 'PART-A',
             'Dest.Storage Bin': 'A1', 'Dest. Storage Type': 'SM', 'Dest.target quantity': 90,
             '_created_at_local': '2026-01-01T08:00:00', '_confirmed_at_local': '2026-01-01T10:00:00'}
    second = dict(first, **{'Transfer Order Number': 'T2', 'Dest.target quantity': 60,
                           '_created_at_local': '2026-01-30T08:00:00', '_confirmed_at_local': '2026-01-30T08:00:00'})
    invalid = dict(first, **{'Transfer Order Number': 'T3', '_created_at_local': '2026-01-30T08:00:00',
                            '_confirmed_at_local': '2026-01-29T08:00:00'})
    return {
        'BIN_STATUS_EXAMPLE': [
            {'Material': 'PART-A', 'Material Description': 'Synthetic bracket', 'Storage Type': 'SM', 'Storage Bin': 'A1', 'Available stock': 10, 'Base Unit of Measure': 'PC'},
            {'Material': 'PART-B', 'Material Description': 'Synthetic bolt', 'Storage Type': 'SM', 'Storage Bin': 'B1', 'Available stock': 0, 'Base Unit of Measure': 'EA'},
        ],
        'HISTORICAL_SM_REPLENISHMENTS': [first, second, deepcopy(first), invalid],
        'HISTORICAL_PROD_DATA': [
            {'TransactionTime': '2026-01-01T10:00:00', 'productionorderOrdNr': 'P1', 'statusstringDescription': 'Passed', 'JobTraceTime (MIN)': 10, 'stationDescr': 'Build'},
            {'TransactionTime': '2026-01-01T11:00:00', 'productionorderOrdNr': 'P1', 'statusstringDescription': 'On hold', 'JobTraceTime (MIN)': 20, 'stationDescr': 'Build'},
        ],
        'PLANNED_PROD_SCHEDULE': [
            {'Sales Order': 'SO1', 'Line Item': '1', '_start_at_local': '2026-02-01T00:00:00'},
            {'Sales Order': 'SO1', 'Line Item': '1', '_start_at_local': '2026-02-01T00:00:00'},
            {'Sales Order': 'SO2', 'Line Item': '1', '_start_at_local': '2026-02-03T00:00:00'},
        ],
    }


def test_real_data_summary_uses_explicit_grains_and_exclusions():
    result = analyze(records(), '1900')
    assert result['inventory_summary']['positions'] == 2
    assert result['inventory_summary']['empty_positions'] == 1
    assert result['production']['events'] == 2
    assert result['production']['distinct_orders'] == 1
    assert result['production']['passed_events'] == 1
    assert result['schedule']['source_rows'] == 3
    assert result['schedule']['unique_lines_by_date'] == 2
    assert result['transfers']['valid_events'] == 2
    assert result['transfers']['duplicate_events'] == 1
    assert result['transfers']['invalid_time_order'] == 1
    assert result['transfers']['median_confirmation_hours'] == 1
    assert result['transfers']['zero_duration_events'] == 1


def test_proxy_uses_transfer_targets_and_never_claims_consumption():
    result = analyze(records(), '1900')
    row = next(r for r in result['inventory'] if r['material'] == 'PART-A')
    assert row['window_transfer_quantity'] == 150
    assert row['window_days'] == 30
    assert row['target_quantity_per_day'] == 5
    assert row['stock_to_target_days'] == 2
    assert result['forecast_validated'] is False
    assert result['snapshot_as_of'] is None
    simulation = scenario(result, row['id'], 2, 0, 8)
    assert simulation['assumed_hours_to_minimum'] == 24
    assert simulation['kind'] == 'assumption_based'
    assert simulation['crossing_at'] is None
    assert simulation['assumptions']


def test_missing_and_ambiguous_inputs_do_not_become_zero_demand():
    data = records()
    data['BIN_STATUS_EXAMPLE'][0]['Available stock'] = None
    result = analyze(data, '1900')
    row = next(r for r in result['inventory'] if r['material'] == 'PART-A')
    assert row['stock'] is None
    assert row['stock_to_target_days'] is None
    with pytest.raises(ValueError): scenario(result, row['id'], 1, 0, 8)
    row = next(r for r in result['inventory'] if r['material'] == 'PART-B')
    with pytest.raises(ValueError): scenario(result, row['id'], 1, 0, 8)


def test_scenario_inputs_and_unsupported_units_rejected():
    data = records()
    data['BIN_STATUS_EXAMPLE'][0]['Base Unit of Measure'] = 'KG'
    result = analyze(data, '1900')
    row = next(r for r in result['inventory'] if r['material'] == 'PART-A')
    with pytest.raises(ValueError): scenario(result, row['id'], 1, 0, 8)
    result = analyze(records(), '1900')
    row = next(r for r in result['inventory'] if r['material'] == 'PART-A')
    with pytest.raises(ValueError): scenario(result, row['id'], -1, 0, 8)


def test_schedule_and_production_do_not_get_joined_into_material_use():
    data = records()
    original = analyze(data, '1900')
    data['PLANNED_PROD_SCHEDULE'] *= 100
    changed = analyze(data, '1900')
    assert original['inventory'] == changed['inventory']
    assert changed['schedule']['source_rows'] == 300


def test_analysis_and_what_if_scenarios_persist_without_altering_source(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    from test_ingest import workbook_bytes
    wb = load_workbook(BytesIO(workbook_bytes()))
    ws = wb['HISTORICAL_SM_REPLENISHMENTS']
    ws['E2'] = 60
    ws['H2'] = ws['F2'].value
    buf = BytesIO(); wb.save(buf)
    path = tmp_path / 'insights.sqlite3'
    store = Store(path)
    ds = store.import_workbook(buf.getvalue(), 'synthetic.xlsx')
    insights = store.insights(ds['id'])
    row = insights['inventory'][0]
    result = store.workbook_scenario(ds['id'], row['id'], 1, 0, 8)
    assert result['assumed_hours_to_minimum'] == 4
    reopened = Store(path)
    assert reopened.insights(ds['id'])['scenarios'][0]['id'] == result['id']
    assert reopened.record_count(ds['id']) == 6
    assert reopened.state(ds['id'])['materials'] == []
    assert reopened.export(ds['id'])['insights']['scenarios'][0]['kind'] == 'assumption_based'


def test_duplicate_stock_is_not_added_twice_and_missing_identifiers_reported():
    data = records()
    data['BIN_STATUS_EXAMPLE'].append(deepcopy(data['BIN_STATUS_EXAMPLE'][0]))
    data['BIN_STATUS_EXAMPLE'].append({'Material': 'MISSING-BIN', 'Available stock': 10})
    result = analyze(data, '1900')
    row = next(r for r in result['inventory'] if r['material'] == 'PART-A')
    assert row['stock'] == 10
    assert result['inventory_summary']['duplicate_rows'] == 1
    assert result['inventory_summary']['missing_key_rows'] == 1


def test_conflicting_duplicate_transfers_and_stock_are_excluded_from_scenarios():
    data = records()
    conflicting = deepcopy(data['HISTORICAL_SM_REPLENISHMENTS'][0])
    conflicting['Dest.target quantity'] = 999
    data['HISTORICAL_SM_REPLENISHMENTS'].append(conflicting)
    result = analyze(data, '1900')
    assert result['transfers']['conflicting_duplicate_keys'] == 1
    assert result['transfers']['valid_events'] == 1
    data['BIN_STATUS_EXAMPLE'].append(dict(data['BIN_STATUS_EXAMPLE'][0], **{'Available stock': 999}))
    result = analyze(data, '1900')
    row = next(r for r in result['inventory'] if r['material'] == 'PART-A')
    assert row['stock'] is None and row['scenario_eligible'] is False


def test_small_positive_target_rate_remains_scenario_eligible():
    data = records()
    for row in data['HISTORICAL_SM_REPLENISHMENTS']:
        row['Dest.target quantity'] = 0.00001
    result = analyze(data, '1900')
    row = next(r for r in result['inventory'] if r['material'] == 'PART-A')
    assert row['scenario_eligible']
    assert row['target_quantity_per_day'] > 0
    saved = scenario(result, row['id'], 1, 0, 8)
    assert saved['assumed_hours_to_minimum'] > 0
